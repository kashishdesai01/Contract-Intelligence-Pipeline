import io
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from hypothesis import given
from hypothesis import strategies as st

from cip.corpus import amendment, base
from cip.domain import analyze, as_of, build_snapshot, extract, parse_bytes, verify_snapshot
from cip.models import Clause, FactQuery, Predicate, digest
from cip.query import lookup, query_facts, search


def snapshot(*texts):
    instruments = [analyze(text, digest(text)) for text in texts]
    return build_snapshot(instruments, {i.id: digest(i.text) for i in instruments})


def test_timeline_and_composed_source_provenance():
    result = snapshot(base(), amendment())
    family = result.families[0]
    assert lookup(result, family.id, "payment_days", date(2025, 3, 1))["value"] == 30
    assert lookup(result, family.id, "payment_days", None)["value"] == 60
    fact = family.versions[-1].facts["payment_days"]
    assert len(fact.evidence) == 2
    assert all(e.quote in result.source_texts[e.source_id] for e in fact.evidence)
    assert not verify_snapshot(result)
    assert as_of(family, date(2024, 1, 1)) is None


def test_sequential_operations_use_evolving_text():
    result = snapshot(
        base(),
        amendment(),
        amendment(
            number=2, when="2025-09-01", operation='In Section 2, replace "60 days" with "90 days".'
        ),
    )
    assert result.families[0].versions[-1].facts["payment_days"].value == 90


def test_same_date_overlap_and_resolution_rebuild():
    texts = [
        base(),
        amendment(),
        amendment(
            number=2, operation='Replace Section 2 with:\n"Invoices are payable within 90 days."'
        ),
    ]
    result = snapshot(*texts)
    latest = result.families[0].versions[-1]
    assert {c.kind for c in latest.conflicts} == {"same_date_overlap"}
    assert all(f.status == "disputed" for f in latest.facts.values())
    assert lookup(result, result.families[0].id, "payment_days", None)["abstained"]
    instruments = [analyze(t, digest(t)) for t in texts]
    rebuilt = build_snapshot(
        instruments, result.source_hashes, resolutions={digest([instruments[2].id, "2"]): "skip"}
    )
    assert rebuilt.families[0].versions[-1].facts["payment_days"].value == 60
    assert digest(rebuilt) != digest(result)


@pytest.mark.parametrize(
    "operation,kind",
    [
        ('In Section 8, replace "30 days" with "60 days".', "unresolved_target"),
        ('In Section 2, replace "500 days" with "60 days".', "text_not_unique"),
        ("Delete Section 2.", "unsupported_operation"),
    ],
)
def test_conflicts_are_never_guessed(operation, kind):
    result = snapshot(base(), amendment(operation=operation))
    assert kind in {c.kind for c in result.families[0].versions[-1].conflicts}
    assert not verify_snapshot(result)


def test_chain_gap_and_retroactive_date_fail_closed():
    result = snapshot(base(), amendment(number=3))
    assert result.families[0].versions[-1].conflicts[0].kind == "incomplete_chain"
    result = snapshot(base(), amendment(when="2024-12-01"))
    assert as_of(result.families[0], date(2024, 12, 15)) is None
    assert result.families[0].versions[-1].conflicts


def test_independent_labeled_example():
    text = Path("tests/fixtures/independent.txt").read_text()
    result = snapshot(text)
    facts = result.families[0].versions[0].facts
    expected = {
        "payment_days": 75,
        "notice_days": 15,
        "auto_renewal": False,
        "renewal_months": 6,
        "governing_law": "Oregon",
        "liability_cap": "unlimited",
        "assignment_consent": False,
        "effective_date": "2025-01-01",
    }
    assert {f: facts[f].value for f in expected} == expected
    assert not verify_snapshot(result)


@pytest.mark.parametrize(
    "text",
    [
        "If approved, Invoices are payable within 30 days.",
        "Invoices are payable within 30 days, unless disputed.",
        "This Agreement shall not automatically renew.",
        "Either party may terminate for convenience with thirty days notice.",
    ],
)
def test_unsupported_semantics_are_unknown(text):
    facts = extract([Clause(key="2", text=text, source_id="source", start=0, end=len(text))])
    assert all(f.status == "unknown" for f in facts.values())


def test_unknown_is_not_false_in_negative_filter():
    result = snapshot(
        base().replace("This Agreement does automatically renew.", "Renewal may be agreed later.")
    )
    query = FactQuery(predicates=[Predicate(field="auto_renewal", op="ne", value=True)])
    receipt = query_facts(result, query)
    assert receipt["count"] == 0
    assert receipt["coverage"]["excluded"] == {"unknown_required_fact": 1}
    assert not receipt["coverage"]["complete"]


def test_aggregate_counts_families_not_evidence_rows():
    result = snapshot(base(), amendment())
    query = FactQuery(predicates=[Predicate(field="payment_days", op="gt", value=45)])
    assert query_facts(result, query)["count"] == 1
    assert (
        query_facts(result, FactQuery(as_of=date(2025, 3, 1), predicates=query.predicates))["count"]
        == 0
    )


def test_reference_precision_distractors_and_ambiguous_base():
    result = snapshot(base(), base("Other"), amendment())
    assert sum(len(f.versions) == 2 for f in result.families) == 1
    result = snapshot(base(), base(payment=45), amendment())
    assert result.review
    assert not query_facts(result, FactQuery())["coverage"]["complete"]


def test_hybrid_retrieval_and_conflict_exclusion():
    result = snapshot(base(), amendment())
    hits = search(result, "Invoices payable days")
    assert hits[0]["clause_key"] == "2"
    assert hits[0]["quote_origin"] == "composite"
    assert all(e["quote"] in result.source_texts[e["source_id"]] for e in hits[0]["evidence"])
    result = snapshot(base(), amendment(operation="Delete Section 2."))
    assert search(result, "Invoices") == []


def test_docx_supported_subset_and_rejection():
    doc = Document()
    for line in base().splitlines():
        doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    assert analyze(
        parse_bytes(
            out.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        "id",
    )
    doc.add_table(1, 1)
    out = io.BytesIO()
    doc.save(out)
    with pytest.raises(ValueError, match="tables"):
        parse_bytes(
            out.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


@given(st.integers(min_value=1, max_value=999), st.integers(min_value=1, max_value=999))
def test_replace_property(old, new):
    result = snapshot(
        base(payment=old),
        amendment(operation=(f'In Section 2, replace "{old} days" with "{new} days".')),
    )
    latest = result.families[0].versions[-1]
    assert latest.facts["payment_days"].value == new
    assert not verify_snapshot(result)
