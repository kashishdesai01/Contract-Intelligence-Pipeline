import asyncio
import io

import pytest
from docx import Document
from docx.oxml import OxmlElement

from cip.api import AskRequest
from cip.corpus import amendment, base
from cip.domain import analyze, build_snapshot, parse_bytes, verify_snapshot
from cip.models import FactQuery, Predicate, digest
from cip.query import lookup, query_facts
from cip.security import BodyLimit


def artifact(*texts):
    instruments = [analyze(text, digest(text)) for text in texts]
    return build_snapshot(instruments, {i.id: digest(i.text) for i in instruments})


def test_ambiguous_amendment_suspends_candidate_families():
    snapshot = artifact(base(), base(payment=45), amendment())
    assert snapshot.review
    assert all(lookup(snapshot, f.id, "payment_days", None)["abstained"] for f in snapshot.families)


def test_preamble_is_not_silently_discarded():
    text = base().replace("\n\n1. Renewal", "\n\nPayment is conditional on approval.\n\n1. Renewal")
    with pytest.raises(ValueError, match="preamble"):
        analyze(text, "source")


def test_different_uploads_of_same_base_do_not_inflate_trusted_counts():
    snapshot = artifact(base(), base().replace("2. Payment", "2. Settlement"))
    result = query_facts(snapshot, FactQuery())
    assert result["count"] == 0
    assert result["coverage"]["excluded"] == {"conflicted": 2}
    assert not result["coverage"]["complete"]


def test_nested_sections_require_explicit_tree_support():
    with pytest.raises(ValueError, match="nested_sections"):
        analyze(base().replace("2. Payment", "2. Payment\n2.1. Detail"), "source")


def test_text_replacement_cannot_renumber_a_section_heading():
    snapshot = artifact(
        base(), amendment(operation='In Section 2, replace "2. Payment" with "7. Payment".')
    )
    assert snapshot.families[0].versions[-1].conflicts
    assert lookup(snapshot, snapshot.families[0].id, "payment_days", None)["abstained"]


def test_accepted_iso_date_forms_match_the_same_agreement():
    snapshot = artifact(base())
    query = FactQuery(predicates=[Predicate(field="effective_date", value="20250101")])
    assert query_facts(snapshot, query)["count"] == 1


def test_unrecognized_qualifier_suspends_semantic_facts():
    text = base().replace(
        "Invoices are payable within 30 days.",
        "Invoices are payable within 30 days.\nSupplier may require prepayment.",
    )
    snapshot = artifact(text)
    assert lookup(snapshot, snapshot.families[0].id, "payment_days", None)["abstained"]


def test_semantic_corruption_detected_by_verifier():
    snapshot = artifact(base())
    snapshot.families[0].versions[0].facts["auto_renewal"].value = False
    assert "invalid_fact_semantics" in verify_snapshot(snapshot)
    snapshot.families[0].versions[0].facts["effective_date"].value = "2030-01-01"
    assert "invalid_date_semantics" in verify_snapshot(snapshot)


def test_review_resolution_must_address_existing_operation():
    snapshot = artifact(base(), amendment())
    instruments = [analyze(base(), digest(base())), analyze(amendment(), digest(amendment()))]
    with pytest.raises(ValueError, match="unknown_review_resolution"):
        build_snapshot(instruments, snapshot.source_hashes, resolutions={"bogus": "skip"})


def test_docx_inherited_numbering_is_rejected():
    doc = Document()
    doc.add_paragraph("An automatically numbered clause", style="List Number")
    output = io.BytesIO()
    doc.save(output)
    with pytest.raises(ValueError, match="automatic_numbering"):
        parse_bytes(
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


def test_docx_unread_footnote_reference_is_not_silently_ignored():
    doc = Document()
    for line in base().splitlines():
        doc.add_paragraph(line)
    doc.paragraphs[5].add_run()._r.append(OxmlElement("w:footnoteReference"))
    output = io.BytesIO()
    doc.save(output)
    with pytest.raises(ValueError, match="unsupported_content"):
        parse_bytes(
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


@pytest.mark.parametrize(
    "field,value,op",
    [
        ("payment_days", True, "eq"),
        ("notice_days", "30", "gt"),
        ("auto_renewal", True, "gt"),
        ("governing_law", 30, "eq"),
    ],
)
def test_filter_types_are_validated(field, value, op):
    with pytest.raises(ValueError):
        Predicate(field=field, value=value, op=op)


def test_chunked_requests_have_enforced_total_size():
    async def scenario():
        messages = iter(
            [
                {"type": "http.request", "body": b"abc", "more_body": True},
                {"type": "http.request", "body": b"def", "more_body": False},
            ]
        )
        sent = []

        async def receive():
            return next(messages)

        async def send(message):
            sent.append(message)

        async def downstream(scope, receive, send):
            await receive()
            await receive()

        await BodyLimit(downstream, limit=5)({"type": "http", "headers": []}, receive, send)
        assert sent[0]["status"] == 413

    asyncio.run(scenario())


@pytest.mark.parametrize(
    "payload",
    [
        {"question": "What are the payment terms?"},
        {"action": "count", "family_id": "one-specific-family"},
        {
            "action": "search",
            "question": "Payment",
            "query": {"predicates": [{"field": "payment_days", "value": 30}]},
        },
    ],
)
def test_assistant_does_not_silently_ignore_question_or_filters(payload):
    with pytest.raises(ValueError):
        AskRequest.model_validate(payload)
