"""Pure, bounded contract semantics. No database, provider, or gold-data dependencies."""

import hashlib
import io
import re
import zipfile
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document

from cip.models import (
    FIELDS,
    Clause,
    Composite,
    Conflict,
    Evidence,
    Fact,
    Family,
    Instrument,
    Operation,
    Snapshot,
    digest,
)

PIPELINE_VERSION = (
    "msa-v1-"
    + hashlib.sha256(
        Path(__file__).read_bytes() + Path(__file__).with_name("models.py").read_bytes()
    ).hexdigest()[:16]
)
PLAYBOOK: dict[str, int | str] = {
    "payment_days": 45,
    "notice_days": 60,
    "liability_cap": "unlimited",
}


class UnsupportedDocument(ValueError):
    pass


def parse_bytes(data: bytes, media: str) -> str:
    if len(data) > 512_000:
        raise UnsupportedDocument("upload_too_large")
    if media == "text/plain":
        text = data.decode("utf-8", errors="strict")
    elif media == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or sum(e.file_size for e in entries) > 2_000_000:
                raise UnsupportedDocument("docx_expansion_limit")
        doc = Document(io.BytesIO(data))

        def numbered(paragraph: Any) -> bool:
            if paragraph._p.xpath(".//w:numPr"):
                return True
            style = paragraph.style
            visited = set()
            while style is not None and style.style_id not in visited:
                visited.add(style.style_id)
                if style.element.xpath(".//w:numPr | .//w:vanish"):
                    return True
                style = style.base_style
            return False

        if doc.tables or any(numbered(p) for p in doc.paragraphs):
            raise UnsupportedDocument("docx_tables_or_automatic_numbering")
        # Fail closed rather than silently ignoring revision or embedded content.
        if doc.element.xpath(
            ".//w:ins | .//w:del | .//w:altChunk | .//w:drawing | .//w:hyperlink | "
            ".//w:object | .//w:vanish | .//w:instrText | .//w:moveFrom | .//w:moveTo | "
            ".//w:sdt | .//w:customXml | .//w:footnoteReference | .//w:endnoteReference"
        ) or any(
            p.text.strip()
            for section in doc.sections
            for part in (section.header, section.footer)
            for p in part.paragraphs
        ):
            raise UnsupportedDocument("docx_unsupported_content")
        if any(
            part.tables or part._element.xpath(".//w:drawing | .//w:sdt")
            for section in doc.sections
            for part in (section.header, section.footer)
        ):
            raise UnsupportedDocument("docx_unsupported_header_footer")
        text = "\n".join(p.text for p in doc.paragraphs)
    else:
        raise UnsupportedDocument("unsupported_media_type")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if not text.strip() or "\x00" in text:
        raise UnsupportedDocument("empty_or_binary")
    return text


def header(text: str, name: str) -> str:
    matches = re.findall(rf"^{re.escape(name)}: (.+)$", text, re.M)
    if len(matches) != 1:
        raise UnsupportedDocument(f"missing_or_duplicate_{name}")
    return matches[0].strip()


def analyze(text: str, source_id: str) -> Instrument:
    effective = date.fromisoformat(header(text, "Effective date"))
    parties = [x.strip() for x in header(text, "Parties").split(" | ")]
    if len(parties) != 2 or any(not p for p in parties) or parties[0] == parties[1]:
        raise UnsupportedDocument("exactly_two_distinct_parties_required")
    first = text.splitlines()[0]
    if first.startswith("Master Services Agreement: "):
        agreement = first.removeprefix("Master Services Agreement: ").strip()
        matches = list(re.finditer(r"^(\d+(?:\.\d+)*)\. (.+)$", text, re.M))
        if not matches or len(matches) > 200:
            raise UnsupportedDocument("requires_1_to_200_numbered_clauses")
        expected_headers = [
            first,
            f"Parties: {' | '.join(parties)}",
            f"Effective date: {effective.isoformat()}",
        ]
        if text.splitlines()[:3] != expected_headers:
            raise UnsupportedDocument("unsupported_header_layout")
        if "\n".join(text[: matches[0].start()].splitlines()[3:]).strip():
            raise UnsupportedDocument("unsupported_preamble")
        clauses = []
        keys = set()
        for i, section in enumerate(matches):
            key = section[1]
            if "." in key:
                raise UnsupportedDocument("nested_sections_unsupported")
            if key in keys:
                raise UnsupportedDocument("duplicate_section")
            keys.add(key)
            start = section.start()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            clauses.append(
                Clause(
                    key=key,
                    text=text[start:end],
                    source_id=source_id,
                    start=start,
                    end=end,
                    history=[source_id],
                    evidence=[Evidence(source_id=source_id, quote=text[start:end])],
                )
            )
        return Instrument(
            id=source_id,
            kind="base",
            agreement=agreement,
            parties=parties,
            effective_date=effective,
            clauses=clauses,
            text=text,
        )
    match = re.fullmatch(r"Amendment (\d+): (.+)", first)
    if not match:
        raise UnsupportedDocument("unsupported_instrument_header")
    operations = []
    # Exact supported grammar: payload is quoted verbatim, never generated by the merger.
    pattern = r'Replace Section ([\d.]+) with:\n"([^"\n]+)"'
    for candidate in re.finditer(pattern, text):
        operations.append(
            Operation(
                kind="replace_clause",
                target=candidate[1],
                new_text=candidate[2],
                source_id=source_id,
                evidence=candidate[0],
                effective_date=effective,
            )
        )
    pattern = r'In Section ([\d.]+), replace "([^"\n]+)" with "([^"\n]+)"\.'
    for candidate in re.finditer(pattern, text):
        operations.append(
            Operation(
                kind="replace_text",
                target=candidate[1],
                old_text=candidate[2],
                new_text=candidate[3],
                source_id=source_id,
                evidence=candidate[0],
                effective_date=effective,
            )
        )
    # Ignore no unknown prose: remove only the recognized header and operation spans.
    if len(operations) > 10:
        raise UnsupportedDocument("operation_count_limit")
    remaining = text
    for op in operations:
        remaining = remaining.replace(op.evidence, "", 1)
    for line in text.splitlines()[:4]:
        remaining = remaining.replace(line, "", 1)
    unsupported = [] if not remaining.strip() and operations else ["unsupported_operation"]
    # Operation order must follow source order, including mixed kinds.
    operations.sort(key=lambda op: text.index(op.evidence))
    return Instrument(
        id=source_id,
        kind="amendment",
        agreement=match[2],
        parties=parties,
        effective_date=effective,
        number=int(match[1]),
        base_date=date.fromisoformat(header(text, "Amends agreement effective")),
        clauses=[],
        text=text,
        operations=operations,
        unsupported=unsupported,
    )


# Full supported sentences only. Number occurrence or heading similarity is never validation.
PATTERNS: dict[str, tuple[str, Any]] = {
    "payment_days": (r"Invoices are payable within (\d{1,3}) days\.", int),
    "notice_days": (
        r"Either party may terminate for convenience with (\d{1,3}) days notice\.",
        int,
    ),
    "auto_renewal": (
        r"This Agreement (does|does not) automatically renew\.",
        lambda v: v == "does",
    ),
    "renewal_months": (r"The renewal term is (\d{1,3}) months\.", int),
    "governing_law": (r"This Agreement is governed by the laws of ([A-Za-z ]+)\.", str),
    "liability_cap": (r"Aggregate liability is (unlimited|limited to annual fees)\.", str),
    "assignment_consent": (
        r"Assignment (requires|does not require) prior written consent\.",
        lambda v: v == "requires",
    ),
}


def extract(clauses: list[Clause], effective: date | None = None) -> dict[str, Fact]:
    facts = {field: Fact(field=field, reason="unsupported_or_absent") for field in FIELDS}
    # The accepted subset is intentionally closed. Extra substantive sentences can qualify
    # otherwise recognizable terms; silently ignoring them would overstate certainty.
    supported = True
    for clause in clauses:
        lines = [line.strip() for line in clause.text.splitlines() if line.strip()]
        if lines and re.fullmatch(r"\d+(?:\.\d+)*\. .+", lines[0]):
            lines = lines[1:]
        if not lines or any(
            not any(re.fullmatch(pattern, line) for pattern, _ in PATTERNS.values())
            for line in lines
        ):
            supported = False
    if not supported:
        return facts
    for field, (pattern, converter) in PATTERNS.items():
        candidates = []
        for clause in clauses:
            # Conditional, exceptional, and referenced clauses require review as a unit.
            if re.search(
                r"\b(if|unless|except|subject to|notwithstanding|Section)\b", clause.text, re.I
            ):
                continue
            for sentence in clause.text.splitlines():
                match = re.fullmatch(pattern, sentence.strip())
                if match:
                    candidates.append(
                        Fact(
                            field=field,
                            value=converter(match[1]),
                            status="disputed" if clause.disputed else "ok",
                            clause_key=clause.key,
                            source_id=clause.source_id,
                            quote=match[0],
                            evidence=clause.evidence,
                        )
                    )
        if len(candidates) == 1:
            facts[field] = candidates[0]
        elif candidates:
            facts[field].reason = "multiple_evidence_candidates"
    if effective:
        # Metadata facts are source-backed separately by the admitted instrument header.
        facts["effective_date"] = Fact(
            field="effective_date",
            value=effective.isoformat(),
            status="ok",
            reason="base_instrument_header",
        )
    return facts


def deviations(facts: dict[str, Fact]) -> list[dict[str, Any]]:
    results = []
    for field, expected in PLAYBOOK.items():
        fact = facts[field]
        if fact.status != "ok":
            continue
        bad = (
            fact.value == expected
            if isinstance(expected, str)
            else (
                isinstance(fact.value, int)
                and not isinstance(fact.value, bool)
                and isinstance(expected, int)
                and fact.value > expected
            )
        )
        if bad:
            results.append(
                {
                    "rule": field,
                    "severity": "high" if field == "liability_cap" else "medium",
                    "value": fact.value,
                    "quote": fact.quote,
                    "source_id": fact.source_id,
                    "playbook": "customer-v1",
                }
            )
    return results


def _conflict(source: str, target: str, kind: str, detail: str) -> Conflict:
    return Conflict(
        id=digest([source, target, kind, detail]),
        source_id=source,
        target=target,
        kind=kind,
        detail=detail,
    )


def build_snapshot(
    instruments: list[Instrument],
    source_hashes: dict[str, str],
    version: str = PIPELINE_VERSION,
    resolutions: dict[str, str] | None = None,
) -> Snapshot:
    """Build deterministically. Reviewer resolutions may skip, never fabricate, an operation."""
    resolutions = resolutions or {}
    known_resolutions = {digest([i.id, op.target]) for i in instruments for op in i.operations}
    if any(key not in known_resolutions or value != "skip" for key, value in resolutions.items()):
        raise ValueError("unknown_review_resolution")
    bases = [i for i in instruments if i.kind == "base"]
    base_identities: dict[tuple[str, tuple[str, ...], date], list[Instrument]] = defaultdict(list)
    for base in bases:
        identity = (
            base.agreement.casefold(),
            tuple(sorted(p.casefold() for p in base.parties)),
            base.effective_date,
        )
        base_identities[identity].append(base)
    duplicate_bases = {b.id for group in base_identities.values() if len(group) > 1 for b in group}
    assignments: dict[str, list[Instrument]] = defaultdict(list)
    review: list[dict[str, str]] = []
    review.extend(
        {"source_id": source, "reason": "ambiguous_base_identity"}
        for source in sorted(duplicate_bases)
    )
    uncertain: dict[str, list[Instrument]] = defaultdict(list)
    for amendment in (i for i in instruments if i.kind == "amendment"):
        candidates = [
            b
            for b in bases
            if b.agreement.casefold() == amendment.agreement.casefold()
            and sorted(p.casefold() for p in b.parties)
            == sorted(p.casefold() for p in amendment.parties)
            and b.effective_date == amendment.base_date
        ]
        if len(candidates) != 1:
            review.append({"source_id": amendment.id, "reason": "ambiguous_or_missing_base"})
            for candidate in candidates:
                uncertain[candidate.id].append(amendment)
        else:
            assignments[candidates[0].id].append(amendment)
    families = []
    for base in sorted(bases, key=lambda b: b.id):
        clauses = [c.model_copy(deep=True) for c in base.clauses]
        conflicts: list[Conflict] = []
        if base.id in duplicate_bases:
            conflicts.append(
                _conflict(
                    base.id,
                    "*",
                    "ambiguous_base_identity",
                    "Multiple base uploads share an agreement identity.",
                )
            )
        versions = []
        amendments = sorted(
            assignments[base.id], key=lambda i: (i.effective_date, i.number or 0, i.id)
        )
        if len(amendments) > 4:
            raise UnsupportedDocument("amendment_chain_limit")
        for amendment in amendments:
            if amendment.effective_date < base.effective_date:
                conflicts.append(
                    _conflict(
                        amendment.id,
                        "*",
                        "out_of_order_reference",
                        "Amendment effective before the base.",
                    )
                )
        dates = sorted(
            {base.effective_date}
            | {a.effective_date for a in amendments if a.effective_date >= base.effective_date}
            | {max(a.effective_date, base.effective_date) for a in uncertain[base.id]}
        )
        seen_numbers: set[int] = set()
        for day in dates:
            for ambiguous in uncertain[base.id]:
                if max(ambiguous.effective_date, base.effective_date) == day:
                    conflicts.append(
                        _conflict(
                            ambiguous.id,
                            "*",
                            "ambiguous_family_link",
                            "An amendment may change this agreement.",
                        )
                    )
            group = [a for a in amendments if a.effective_date == day]
            targets: dict[str, set[str]] = defaultdict(set)
            for a in group:
                for op in a.operations:
                    if resolutions.get(digest([a.id, op.target])) != "skip":
                        targets[op.target].add(a.id)
            for amendment in group:
                if amendment.number is not None:
                    if (
                        amendment.number in seen_numbers
                        or amendment.number != len(seen_numbers) + 1
                    ):
                        conflicts.append(
                            _conflict(
                                amendment.id,
                                "*",
                                "incomplete_chain",
                                "Amendment numbering is not contiguous.",
                            )
                        )
                    seen_numbers.add(amendment.number)
                if day < base.effective_date or amendment.unsupported:
                    conflicts.append(
                        _conflict(
                            amendment.id,
                            "*",
                            "unsupported_operation",
                            "Unsupported operation or date order.",
                        )
                    )
                for op in amendment.operations:
                    if resolutions.get(digest([amendment.id, op.target])) == "skip":
                        continue
                    if len(targets[op.target]) > 1:
                        conflicts.append(
                            _conflict(
                                amendment.id,
                                op.target,
                                "same_date_overlap",
                                "Multiple instruments change this section.",
                            )
                        )
                        continue
                    selected = [c for c in clauses if c.key == op.target]
                    if len(selected) != 1:
                        conflicts.append(
                            _conflict(
                                amendment.id,
                                op.target,
                                "unresolved_target",
                                "Target section does not exist uniquely.",
                            )
                        )
                        continue
                    clause = selected[0]
                    if clause.disputed:
                        conflicts.append(
                            _conflict(
                                amendment.id,
                                op.target,
                                "disputed_dependency",
                                "Prior conflict requires review.",
                            )
                        )
                        continue
                    if op.kind == "replace_text":
                        first_line = clause.text.splitlines()[0]
                        body_start = (
                            min(len(clause.text), len(first_line) + 1)
                            if re.fullmatch(r"\d+\. .+", first_line)
                            else 0
                        )
                        body = clause.text[body_start:]
                        if not op.old_text or body.count(op.old_text) != 1:
                            conflicts.append(
                                _conflict(
                                    amendment.id,
                                    op.target,
                                    "text_not_unique",
                                    "Old text must occur exactly once.",
                                )
                            )
                            continue
                        clause.text = clause.text[:body_start] + body.replace(
                            op.old_text, op.new_text, 1
                        )
                    else:
                        clause.text = op.new_text + "\n"
                    clause.source_id = amendment.id
                    clause.history.append(amendment.id)
                    if op.kind == "replace_clause":
                        clause.evidence = [Evidence(source_id=amendment.id, quote=op.new_text)]
                    else:
                        clause.evidence.append(Evidence(source_id=amendment.id, quote=op.evidence))
                    # Composite offsets are not advertised as offsets into the amendment.
                    clause.start, clause.end = 0, len(clause.text)
            if conflicts:
                # Conservative dependency closure: a conflict can alter definitions anywhere.
                for clause in clauses:
                    clause.disputed = True
            facts = extract(clauses, base.effective_date)
            facts["effective_date"].source_id = base.id
            facts["effective_date"].quote_origin = "source_header"
            facts["effective_date"].quote = f"Effective date: {base.effective_date.isoformat()}"
            facts["effective_date"].evidence = [
                Evidence(
                    source_id=base.id, quote=f"Effective date: {base.effective_date.isoformat()}"
                )
            ]
            if conflicts:
                for fact in facts.values():
                    fact.status, fact.reason = "disputed", "unresolved_composite_conflict"
            versions.append(
                Composite(
                    as_of=day,
                    clauses=[c.model_copy(deep=True) for c in clauses],
                    facts=facts,
                    conflicts=list(conflicts),
                    deviations=deviations(facts),
                )
            )
        families.append(
            Family(
                id=base.id,
                agreement=base.agreement,
                parties=base.parties,
                base_id=base.id,
                versions=versions,
            )
        )
    return Snapshot(
        version=version,
        families=families,
        review=review,
        source_hashes=source_hashes,
        source_texts={i.id: i.text for i in instruments},
    )


def as_of(family: Family, when: date | None) -> Composite | None:
    versions = [v for v in family.versions if when is None or v.as_of <= when]
    return max(versions, key=lambda v: v.as_of) if versions else None


def verify_snapshot(snapshot: Snapshot) -> list[str]:
    errors = []
    for family in snapshot.families:
        if family.base_id not in snapshot.source_hashes:
            errors.append("missing_base_source")
        for version in family.versions:
            expected_facts = extract(version.clauses)
            for clause in version.clauses:
                if any(source not in snapshot.source_hashes for source in clause.history):
                    errors.append("missing_provenance")
            for fact in version.facts.values():
                if fact.status == "ok" and fact.field == "effective_date":
                    try:
                        actual_date = header(
                            snapshot.source_texts[family.base_id], "Effective date"
                        )
                        if fact.value != date.fromisoformat(actual_date).isoformat():
                            errors.append("invalid_date_semantics")
                    except (KeyError, ValueError):
                        errors.append("invalid_date_semantics")
                if (
                    fact.status == "ok"
                    and fact.field != "effective_date"
                    and (
                        expected_facts[fact.field].status != "ok"
                        or expected_facts[fact.field].value != fact.value
                    )
                ):
                    errors.append("invalid_fact_semantics")
                if (
                    fact.status == "ok"
                    and fact.quote
                    and fact.field != "effective_date"
                    and not any(
                        c.key == fact.clause_key
                        and c.source_id == fact.source_id
                        and fact.quote in c.text
                        and not c.disputed
                        for c in version.clauses
                    )
                ):
                    errors.append("invalid_fact_evidence")
                if fact.status == "ok" and (
                    not fact.evidence
                    or any(
                        e.quote not in snapshot.source_texts.get(e.source_id, "")
                        for e in fact.evidence
                    )
                ):
                    errors.append("invalid_source_evidence")
            if version.conflicts and any(f.status == "ok" for f in version.facts.values()):
                errors.append("trusted_conflicted_fact")
    return errors
